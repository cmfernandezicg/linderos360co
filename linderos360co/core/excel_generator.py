# -*- coding: utf-8 -*-
"""
excel_generator.py
==================
Módulo de exportación de coordenadas de vértices a formato .xlsx.

Hoja "Vértices":
    N° | Marca | Norte (m) | Este (m) | Latitud (°) | Longitud (°) | Distancia (m) | Azimut (°)

Precisión:
    - Norte / Este  : 4 decimales  (EPSG:9377, consistente con módulo TXT)
    - Latitud / Lon : 6 decimales  (EPSG:4326, ~0.1 m catastral)
    - Distancia     : 1 decimal    (decímetro, según normativa IGAC)
    - Azimut        : 4 decimales  (grados decimales)

El orden de vértices es normativo: inicio en vértice noroccidental, sentido horario.
La última fila de cada anillo (punto de cierre) no incluye Distancia ni Azimut.
La fila de totales al final de cada sección muestra conteo y suma de distancias.

Soporte de anillos interiores
------------------------------
Los anillos interiores se escriben en la misma hoja a continuación del exterior,
separados por una fila de encabezado gris con el rótulo "ANILLO INTERIOR N°X".
La numeración de vértices (N°) se reinicia en 1 para cada anillo.
Las marcas del exterior son P1, P2... Las del interior N°1 son PI1-1, PI1-2...

Lógica de nombre de archivo
---------------------------
Cuando el módulo TXT está activo junto con Excel, se reutilizan los campos
"campo_nupre" y "campo_fmi" ya mapeados — el usuario no necesita mapear dos veces.
Cuando el módulo TXT está desactivado, se usan "campo_nupre_xlsx" / "campo_fmi_xlsx"
propios del tab Excel. Fallback final: FID del feature.
"""

import os
from typing import Optional

from qgis.core import (
    QgsFeature,
    QgsCoordinateReferenceSystem,
    QgsCoordinateTransform,
    QgsProject,
)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

from ..core.geometry_utils import (
    reproyectar_geometry,
    obtener_vertices,
    asegurar_sentido_horario,
    punto_noroccidental,
    reordenar_desde_inicio,
    calcular_azimut,
    calcular_distancia,
    preparar_anillos_interiores_normativos,
)

# ---------------------------------------------------------------------------
# Constantes de precisión
# ---------------------------------------------------------------------------
_DECIMALES_PLANAS = 4   # Norte / Este (EPSG:9377)
_DECIMALES_GEO    = 6   # Latitud / Longitud (EPSG:4326)
_DECIMALES_DIST   = 1   # Distancia (metros, 1 decimal)
_DECIMALES_AZIMUT = 4   # Azimut decimal

_CRS_9377 = QgsCoordinateReferenceSystem("EPSG:9377")
_CRS_4326 = QgsCoordinateReferenceSystem("EPSG:4326")

_NULOS = {"", "NULL", "None", "null", "nan", "NaN", "none"}

# ---------------------------------------------------------------------------
# Estilos visuales
# ---------------------------------------------------------------------------
_COLOR_ENCABEZADO        = "1F4E79"   # azul institucional oscuro
_COLOR_FILA_PAR          = "DDEEFF"   # azul claro para filas alternas
_COLOR_SEP_ANILLO        = "D9D9D9"   # gris para fila separadora de anillo interior
_COLOR_SEP_ANILLO_FUENTE = "1F4E79"   # texto azul oscuro en separador


def _estilos() -> tuple:
    """Construye y retorna los objetos de estilo openpyxl."""
    fuente_enc  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    relleno_enc = PatternFill("solid", fgColor=_COLOR_ENCABEZADO)
    al_centro   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    al_derecha  = Alignment(horizontal="right",  vertical="center")
    borde_fino  = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    return fuente_enc, relleno_enc, al_centro, al_derecha, borde_fino


# ---------------------------------------------------------------------------
# Helpers internos
# ---------------------------------------------------------------------------

def _nombre_archivo(feature: QgsFeature, config: dict) -> str:
    """
    Determina la ruta del archivo de salida.

    Prioridad para el nombre base:
      1. campo_nupre       → módulo TXT activo (reutilización directa)
      2. campo_nupre_xlsx  → módulo TXT desactivado, campo propio del tab Excel
      3. campo_fmi         → módulo TXT activo
      4. campo_fmi_xlsx    → módulo TXT desactivado
      5. FID del feature   → último recurso

    Si el archivo ya existe agrega sufijo _1, _2...
    """
    ruta_salida = config["ruta_salida"]

    def _val_campo(campo_key: str) -> Optional[str]:
        campo = config.get(campo_key)
        if not campo:
            return None
        try:
            val = str(feature[campo]).strip()
        except Exception:
            return None
        return val if val not in _NULOS else None

    base = (
        _val_campo("campo_nupre")
        or _val_campo("campo_nupre_xlsx")
        or _val_campo("campo_fmi")
        or _val_campo("campo_fmi_xlsx")
        or str(feature.id())
    )
    base = "".join(c for c in base if c.isalnum() or c in "._- ").strip()
    base = base or str(feature.id())

    ruta = os.path.join(ruta_salida, f"{base}.xlsx")
    contador = 1
    while os.path.exists(ruta):
        ruta = os.path.join(ruta_salida, f"{base}_{contador}.xlsx")
        contador += 1
    return ruta


def _punto_a_wgs84(punto, transform: QgsCoordinateTransform) -> tuple:
    """
    Reproyecta un QgsPointXY de EPSG:9377 a EPSG:4326.
    Retorna (latitud, longitud) en grados decimales.
    """
    p = transform.transform(punto)
    return p.y(), p.x()


# ---------------------------------------------------------------------------
# Construcción de filas por anillo
# ---------------------------------------------------------------------------

def _construir_filas_anillo(vertices: list, transform: QgsCoordinateTransform,
                             prefijo_marca: str = "P") -> list:
    """
    Genera la lista de tuplas para un anillo (exterior o interior).

    Cada tupla: (N°, Marca, Norte, Este, Latitud, Longitud, Distancia, Azimut)

    La última fila lleva None en Distancia y Azimut: el recorrido normativo
    cierra en P1 y ese segmento ya fue descrito en la fila anterior.

    Norte = coordenada Y en EPSG:9377 (Northing)
    Este  = coordenada X en EPSG:9377 (Easting)
    """
    filas = []
    n = len(vertices)

    for i, punto in enumerate(vertices):
        norte = round(punto.y(), _DECIMALES_PLANAS)
        este  = round(punto.x(), _DECIMALES_PLANAS)
        lat, lon = _punto_a_wgs84(punto, transform)
        lat = round(lat, _DECIMALES_GEO)
        lon = round(lon, _DECIMALES_GEO)

        if i < n - 1:
            p_sig  = vertices[i + 1]
            dist   = round(calcular_distancia(punto, p_sig), _DECIMALES_DIST)
            azimut = round(calcular_azimut(punto, p_sig), _DECIMALES_AZIMUT)
        else:
            # Último vértice: cierra el recorrido, sin segmento nuevo
            dist   = None
            azimut = None

        marca = f"{prefijo_marca}{i + 1}"
        filas.append((i + 1, marca, norte, este, lat, lon, dist, azimut))

    return filas


# ---------------------------------------------------------------------------
# Escritura de la hoja Excel
# ---------------------------------------------------------------------------

def _escribir_fila_separador(ws, fila_idx: int, texto: str,
                              borde, n_columnas: int = 8) -> None:
    """Escribe una fila de separador con fondo gris para distinguir anillos interiores."""
    relleno_sep = PatternFill("solid", fgColor=_COLOR_SEP_ANILLO)
    fuente_sep  = Font(name="Calibri", bold=True, italic=True,
                       color=_COLOR_SEP_ANILLO_FUENTE, size=10)
    al_centro   = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, n_columnas + 1):
        celda        = ws.cell(row=fila_idx, column=col_idx)
        celda.fill   = relleno_sep
        celda.border = borde
        celda.font   = fuente_sep
        celda.alignment = al_centro

    # Texto solo en columna 2 (Marca), que es la columna más legible
    ws.cell(row=fila_idx, column=2, value=texto)


def _escribir_fila_total(ws, fila_idx: int, n_vertices: int, total_dist: float,
                         fuente_total, borde, al_centro, al_derecha) -> None:
    """Escribe la fila de totales para un anillo."""
    celda_lbl           = ws.cell(row=fila_idx, column=2,
                                  value=f"Total: {n_vertices} vértices")
    celda_lbl.font      = fuente_total
    celda_lbl.border    = borde
    celda_lbl.alignment = al_centro

    celda_dist               = ws.cell(row=fila_idx, column=7, value=total_dist)
    celda_dist.font          = fuente_total
    celda_dist.border        = borde
    celda_dist.alignment     = al_derecha
    celda_dist.number_format = "#,##0.0"

    # Borde en celdas vacías de la fila de totales
    for col_idx in [1, 3, 4, 5, 6, 8]:
        ws.cell(row=fila_idx, column=col_idx).border = borde


def _escribir_filas_datos(ws, filas: list, fila_inicio: int,
                          fuente_dato, borde, al_centro, al_derecha) -> int:
    """
    Escribe las filas de datos de un anillo en la hoja.
    Retorna el índice de la siguiente fila disponible.
    """
    relleno_par = PatternFill("solid", fgColor=_COLOR_FILA_PAR)

    for offset, fila in enumerate(filas):
        fila_idx = fila_inicio + offset
        es_par   = (fila_idx % 2 == 0)

        for col_idx, valor in enumerate(fila, start=1):
            celda           = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.font      = fuente_dato
            celda.border    = borde
            celda.alignment = al_centro if col_idx <= 2 else al_derecha
            if es_par:
                celda.fill = relleno_par

            # Formato numérico por columna
            if col_idx in (3, 4):    # Norte / Este
                celda.number_format = "#,##0.0000"
            elif col_idx in (5, 6):  # Latitud / Longitud
                celda.number_format = "0.000000"
            elif col_idx == 7:       # Distancia
                celda.number_format = "#,##0.0"
            elif col_idx == 8:       # Azimut
                celda.number_format = "0.0000"

    return fila_inicio + len(filas)


def _escribir_hoja_vertices(wb, filas_exterior: list,
                             filas_interiores: list) -> None:
    """
    Escribe, formatea y congela la hoja Vértices en el workbook.

    filas_exterior  : lista de tuplas para el anillo exterior
    filas_interiores: lista de listas de tuplas (una lista por anillo interior)
    """
    fuente_enc, relleno_enc, al_centro, al_derecha, borde = _estilos()
    fuente_dato  = Font(name="Calibri", size=10)
    fuente_total = Font(name="Calibri", bold=True, size=10)

    ws = wb.active
    ws.title = "Vértices"

    # --- Encabezados ---
    encabezados = [
        "N°", "Marca",
        "Norte (m)", "Este (m)",
        "Latitud (°)", "Longitud (°)",
        "Distancia (m)", "Azimut (°)",
    ]
    anchos_col = [6, 9, 16, 16, 14, 14, 15, 13]

    for col_idx, (texto, ancho) in enumerate(zip(encabezados, anchos_col), start=1):
        celda           = ws.cell(row=1, column=col_idx, value=texto)
        celda.font      = fuente_enc
        celda.fill      = relleno_enc
        celda.alignment = al_centro
        celda.border    = borde
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    ws.row_dimensions[1].height = 28

    # --- Anillo exterior ---
    fila_actual = 2
    fila_actual = _escribir_filas_datos(
        ws, filas_exterior, fila_actual,
        fuente_dato, borde, al_centro, al_derecha
    )

    total_dist_ext = round(
        sum(f[6] for f in filas_exterior if f[6] is not None), _DECIMALES_DIST
    )
    _escribir_fila_total(
        ws, fila_actual, len(filas_exterior), total_dist_ext,
        fuente_total, borde, al_centro, al_derecha
    )
    fila_actual += 1

    # --- Anillos interiores ---
    for idx_anillo, filas_int in enumerate(filas_interiores, start=1):
        # Fila separadora
        _escribir_fila_separador(
            ws, fila_actual,
            f"ANILLO INTERIOR N°{idx_anillo}",
            borde
        )
        fila_actual += 1

        # Filas de datos del anillo interior
        fila_actual = _escribir_filas_datos(
            ws, filas_int, fila_actual,
            fuente_dato, borde, al_centro, al_derecha
        )

        # Fila de totales del anillo interior
        total_dist_int = round(
            sum(f[6] for f in filas_int if f[6] is not None), _DECIMALES_DIST
        )
        _escribir_fila_total(
            ws, fila_actual, len(filas_int), total_dist_int,
            fuente_total, borde, al_centro, al_derecha
        )
        fila_actual += 1

    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# API pública
# ---------------------------------------------------------------------------

def exportar_xlsx(feature: QgsFeature, config: dict) -> str:
    """
    Genera el archivo .xlsx de coordenadas de vértices para el feature dado.

    Soporta predios con anillos interiores. Los anillos interiores se añaden
    en la misma hoja bajo filas separadoras grises "ANILLO INTERIOR N°X".

    El pipeline geométrico es idéntico al módulo TXT:
    reproyección → vértices → sentido horario → inicio noroccidental.

    Parámetros
    ----------
    feature : QgsFeature
        Feature del predio seleccionado.
    config : dict
        Config dict del plugin. Claves relevantes para este módulo:
            - "capa"             : QgsVectorLayer
            - "ruta_salida"      : str
            - "campo_nupre"      : str | None  (del módulo TXT si está activo)
            - "campo_fmi"        : str | None  (del módulo TXT si está activo)
            - "campo_nupre_xlsx" : str | None  (propio, cuando TXT está desactivado)
            - "campo_fmi_xlsx"   : str | None  (propio, cuando TXT está desactivado)

    Retorna
    -------
    str
        Ruta absoluta del archivo .xlsx generado.

    Lanza
    -----
    ImportError
        Si openpyxl no está instalado en el entorno QGIS.
    Exception
        Si falla la escritura del archivo o el pipeline geométrico.
    """
    if not OPENPYXL_OK:
        raise ImportError(
            "openpyxl no está disponible en este entorno QGIS. "
            "Instálalo con: pip install openpyxl --break-system-packages"
        )

    capa = config["capa"]

    # Pipeline geométrico — anillo exterior (idéntico a v1.0.1)
    geom      = reproyectar_geometry(feature.geometry(), capa.crs())
    verts_raw = obtener_vertices(geom)
    verts_cw  = asegurar_sentido_horario(verts_raw)
    idx_nw    = punto_noroccidental(verts_cw)
    vertices  = reordenar_desde_inicio(verts_cw, idx_nw)

    # Pipeline geométrico — anillos interiores
    anillos_interiores = preparar_anillos_interiores_normativos(
        feature.geometry(), capa.crs()
    )

    # Transformación EPSG:9377 → EPSG:4326 para columnas geográficas
    transform = QgsCoordinateTransform(_CRS_9377, _CRS_4326, QgsProject.instance())

    # Construir filas para cada anillo
    filas_exterior = _construir_filas_anillo(vertices, transform, prefijo_marca="P")

    filas_interiores = []
    for idx_anillo, verts_int in enumerate(anillos_interiores, start=1):
        prefijo = f"PI{idx_anillo}-"
        filas_interiores.append(
            _construir_filas_anillo(verts_int, transform, prefijo_marca=prefijo)
        )

    wb = openpyxl.Workbook()
    _escribir_hoja_vertices(wb, filas_exterior, filas_interiores)

    ruta = _nombre_archivo(feature, config)
    wb.save(ruta)
    return ruta
